"""提取「表情插件用法」docx 的文本"""
from pathlib import Path

from docx import Document

SRC = Path("人格提示词/表情插件用法(1).docx")
OUT = Path("人格提示词/表情插件用法_提取.txt")

doc = Document(SRC)
lines: list[str] = []

for idx, para in enumerate(doc.paragraphs, 1):
    text = para.text.strip()
    if text:
        lines.append(f"[{idx:03d}][{para.style.name}] {text}")

for t_idx, table in enumerate(doc.tables, 1):
    lines.append(f"\n=== 表格 {t_idx} ===")
    for row in table.rows:
        lines.append(" | ".join(cell.text.strip() for cell in row.cells))

OUT.write_text("\n".join(lines), encoding="utf-8")
print(f"段落数: {len(doc.paragraphs)}, 表格数: {len(doc.tables)}, 输出 {OUT.stat().st_size} 字节")
